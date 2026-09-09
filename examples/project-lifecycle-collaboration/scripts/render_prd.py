"""Create an editable draft in the caller's local directory; no network I/O."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from contracts import (ContractError, check_context, index_by, load_json,
                       require_same, validate)


def validate_document(context: dict[str, Any], payload: dict[str, Any]) -> dict[str, Any]:
    check_context(context)
    validate("document", payload)
    require_same(context, payload, ("project_id", "context_revision"))
    evidence = index_by(payload["evidence"], "evidence_id")
    trusted = index_by(context["evidence"], "evidence_id")
    if any(trusted.get(key) != item for key, item in evidence.items()):
        raise ContractError("document_evidence_changed")
    for section in payload["sections"]:
        refs = [ref for paragraph in section["paragraphs"] for ref in paragraph["evidence_ids"]]
        if any(ref not in evidence for ref in refs):
            raise ContractError("unknown_document_evidence")
    return evidence


def render_prd(context: dict[str, Any], payload: dict[str, Any], output: Path) -> dict[str, Any]:
    evidence = validate_document(context, payload)
    if output.suffix.lower() != ".docx":
        raise ContractError("docx_extension_required")
    if output.exists():
        raise ContractError("output_already_exists")
    document = Document()
    for tree in (document.element, document.styles.element):
        for border in tree.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    page = document.sections[0]
    page.page_width, page.page_height = Cm(21.59), Cm(27.94)
    page.top_margin = page.bottom_margin = Cm(2)
    page.left_margin = page.right_margin = Cm(2.3)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor.from_string("000000")
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15
    document.styles["Title"].font.size = Pt(24)
    document.styles["Title"].paragraph_format.space_after = Pt(10)
    for name, size in (("Heading 1", 14), ("Heading 2", 11)):
        document.styles[name].font.size = Pt(size)
        document.styles[name].paragraph_format.space_before = Pt(10)
        document.styles[name].paragraph_format.space_after = Pt(4)
    document.add_heading(payload["title"], 0)
    document.add_paragraph(f"版本 {payload['version']} | 待评审草稿")
    document.add_paragraph(f"项目 {payload['project_id']} | 上下文版本 {payload['context_revision']}")
    document.add_paragraph("本文为需求草稿；正式范围、研发估算及交付承诺以授权人员对对应内容版本的评审决定为准。")
    quoted_evidence: set[str] = set()
    for section in payload["sections"]:
        heading = re.sub(r"^\s*(?:\d+|[一二三四五六七八九十]+)[、.．]\s*", "", section["heading"])
        document.add_heading(heading, level=1)
        for paragraph in section["paragraphs"]:
            if paragraph["kind"] == "evidence":
                for ref in paragraph["evidence_ids"]:
                    item = evidence[ref]
                    quoted_evidence.add(ref)
                    state = "已核验记录" if item["status"] == "verified" else "来源自述"
                    document.add_paragraph(f"{ref} {state}：{item['text']}")
            else:
                prefix = "建议：" if paragraph["kind"] == "suggestion" else "待确认："
                text = paragraph["text"]
                if paragraph["kind"] == "suggestion":
                    text = re.sub(r"^建议[：:\s]*", "", text)
                document.add_paragraph(prefix + text)
        refs = list(dict.fromkeys(ref for paragraph in section["paragraphs"] for ref in paragraph["evidence_ids"]))
        if refs:
            paragraph = document.add_paragraph("依据：" + "、".join(refs))
            paragraph.runs[0].font.size = Pt(9)
    for heading, values in (("待确认事项", payload["open_questions"]),
                            ("评审备注草稿", payload["review_notes"])):
        document.add_heading(heading, level=1)
        for text in values or ["本次未提供。"]:
            document.add_paragraph(("待核对：" if heading == "评审备注草稿" else "") + text)
    document.add_heading("证据目录", level=1)
    for item in evidence.values():
        state = "已核验" if item["status"] == "verified" else "来源自述，待核验"
        document.add_paragraph(f"{item['evidence_id']} {state}；来源引用：{item['source_ref']}")
        if item["evidence_id"] not in quoted_evidence:
            document.add_paragraph(item["text"])
    footer = page.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run(f"{payload['version']} | 第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    # Exclusive creation prevents a concurrent export from overwriting a file.
    with output.open("xb") as stream:
        document.save(stream)
    digest = hashlib.sha256(output.read_bytes()).hexdigest()
    return {"status": "draft_created", "sha256": digest,
            "bytes": output.stat().st_size, "delivered": False, "approved": False}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("context")
    parser.add_argument("document")
    parser.add_argument("output")
    args = parser.parse_args()
    try:
        result = render_prd(load_json(args.context), load_json(args.document), Path(args.output))
    except (ContractError, OSError) as exc:
        print(json.dumps({"status": "failed", "code": str(exc) if isinstance(exc, ContractError) else "file_unavailable"}))
        return 2
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
