"""Cell evidence, filtering, and exclusion; never substitutes for live acceptance."""
from openpyxl import Workbook
from io import BytesIO
from zipfile import ZipFile
import pytest

from app.distillation_sample_schemas import DatabaseSampleArguments
from app.services.library_spreadsheet_reader import inspect_path
from app.services.library_xlsx_stream import BoundedXML, rows


def test_file_samples_include_actual_cells_row_numbers_and_exact_filters(tmp_path):
    path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["case", "period", "result", "password"])
    sheet.append(["r1", "2026-01", "closed", "synthetic-secret"])
    sheet.append(["r1", "2026-02", "reworked", "synthetic-secret"])
    workbook.save(path)
    args = DatabaseSampleArguments(data_source_id="source", bucket_file_id="file", limit=1)
    first = inspect_path(path, args)
    fields = {item["name"]: item["field_key"] for item in first["table"]["fields"]}
    assert "password" not in fields
    assert first["rows"] == [{fields["case"]: "r1", fields["period"]: "2026-01", fields["result"]: "closed"}]
    assert first["source_rows"] == [2]
    assert first["has_more"] is True
    filtered = inspect_path(path, args.model_copy(update={"table_key": first["table"]["table_key"],
        "filters": [args.__class__.model_validate({"data_source_id": "source",
            "filters": [{"field_key": fields["period"], "value": "2026-02"}]}).filters[0]]}))
    assert filtered["source_rows"] == [3]
    assert filtered["rows"][0][fields["result"]] == "reworked"


def test_empty_template_is_not_a_historical_result_and_unknown_handles_fail(tmp_path):
    path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    workbook.active.append(["id", "result"])
    workbook.save(path)
    args = DatabaseSampleArguments(data_source_id="source", bucket_file_id="file")
    sample = inspect_path(path, args)
    assert sample["rows"] == []
    assert sample["source_rows"] == []
    with pytest.raises(ValueError, match="引用"):
        inspect_path(path, args.model_copy(update={"table_key": "0" * 32}))


def test_multiple_sheets_require_explicit_selection_and_large_cells_are_excluded(tmp_path):
    path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    workbook.active.append(["id", "text"])
    workbook.active.append(["one", "x" * 1100])
    workbook.create_sheet("other").append(["id", "result"])
    workbook.save(path)
    args = DatabaseSampleArguments(data_source_id="source", bucket_file_id="file")
    directory = inspect_path(path, args)
    assert directory["kind"] == "database_catalog"
    table = directory["tables"][0]
    sample = inspect_path(path, args.model_copy(update={"table_key": table["table_key"]}))
    assert sample["rows"][0][table["fields"][1]["field_key"]] is None
    assert len(sample["excluded_cells"]) == 1


def test_xml_rejects_entities_even_when_marker_crosses_read_boundary():
    reader = BoundedXML(BytesIO(b'<!DOC' + b'TYPE root [<!ENTITY x "text">]><root/>'))
    reader.read(5)
    with pytest.raises(ValueError, match="XML"):
        reader.read(100)


def test_shared_text_is_selected_by_reference_and_preserves_sparse_cells(tmp_path):
    from datetime import datetime

    path = tmp_path / "shared.xlsx"
    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    with ZipFile(path, "w") as archive:
        archive.writestr("sheet.xml", f'<worksheet xmlns="{ns}"><sheetData><row r="7"><c r="A7" t="s"><v>1</v></c><c r="C7" t="s"><v>0</v></c></row></sheetData></worksheet>')
        archive.writestr("xl/sharedStrings.xml", f'<sst xmlns="{ns}"><si><t>result</t></si><si><r><t>case</t></r><r><t>-one</t></r></si></sst>')
    with ZipFile(path) as archive:
        assert rows(archive, "sheet.xml", 1, set(), datetime(1899, 12, 30)) == [(7, ["case-one", None, "result"])]


def test_sparse_formatting_does_not_expand_data_and_wide_business_tables_are_bounded(tmp_path):
    from datetime import datetime

    path = tmp_path / "wide.xlsx"
    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    with ZipFile(path, "w") as archive:
        archive.writestr("sheet.xml", f'<worksheet xmlns="{ns}"><sheetData><row r="1"><c r="CC1"><v>42</v></c><c r="XFD1" s="0"/></row></sheetData></worksheet>')
    with ZipFile(path) as archive:
        result = rows(archive, "sheet.xml", 1, set(), datetime(1899, 12, 30))
    assert len(result[0][1]) == 81
    assert result[0][1][-1] == "42"


def test_title_rows_are_not_misreported_as_single_column_business_data(tmp_path):
    path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Document title"])
    sheet.merge_cells("A1:C1")
    sheet.append(["id", "rule", "basis"])
    sheet.append(["one", "Check evidence", "Written policy"])
    workbook.save(path)
    sample = inspect_path(path, DatabaseSampleArguments(data_source_id="source", bucket_file_id="file"))
    fields = {field["name"]: field["field_key"] for field in sample["table"]["fields"]}
    assert sample["header_row"] == 2
    assert sample["source_rows"] == [3]
    assert sample["rows"][0][fields["rule"]] == "Check evidence"
