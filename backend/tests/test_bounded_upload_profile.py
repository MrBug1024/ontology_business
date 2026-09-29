from openpyxl import Workbook
from datetime import datetime
from zipfile import ZipFile

from app.services import catalog_ingestion_service as ingestion
from app.services import input_contract_schema


def test_excel_structure_discovery_does_not_scan_all_business_rows():
    consumed = 0

    def rows():
        nonlocal consumed
        yield ("id", "amount")
        for index in range(100_000):
            consumed += 1
            yield (index, 2.5)

    _, header, sample, truncated, count = ingestion._sheet_profile_rows(rows())
    assert header == ["id", "amount"]
    assert consumed <= ingestion.MAX_PROFILE_ROWS + 26
    assert len(sample) == ingestion.MAX_PROFILE_ROWS
    assert truncated is True
    assert count is None


def test_excel_small_sheet_retains_exact_count():
    _, _, sample, truncated, count = ingestion._sheet_profile_rows([
        ("id", "amount"), ("x", 2.5), ("y", 3.5),
    ])
    assert len(sample) == count == 2
    assert truncated is False


def test_csv_profile_marks_unknown_total_count(tmp_path):
    path = tmp_path / "large.csv"
    path.write_text("id,amount\n" + "x,2.5\n" * 100_000, encoding="utf-8")
    profile = ingestion._csv_profile_path(path, ingestion._FORMAT_SPECS[".csv"])
    table = profile["tables"][0]
    assert table["record_count"] is None
    assert table["record_count_exact"] is False
    assert table["sample_row_count"] == ingestion.MAX_PROFILE_ROWS
    assert table["sample_truncated"] is True
    assert len(input_contract_schema.structural_fingerprint(profile)) == 64

    from app.services.input_contract_validator import validate_profile
    from app.services.input_contract_schema import CONTENT_CONTRACT_KEY, CONTENT_CONTRACT_VERSION, InputContractError
    import pytest
    relation = {"fields": [{"name": "id", "logical_type": "string", "required": True}],
                "allow_additional_fields": True, "minimum_data_rows": 1}
    schema = {CONTENT_CONTRACT_KEY: {"version": CONTENT_CONTRACT_VERSION,
                                    "relations": [relation], "allow_additional_relations": False}}
    assert validate_profile(schema, profile).matched
    relation['minimum_data_rows'] = 1001
    with pytest.raises(InputContractError, match='structural contract'):
        validate_profile(schema, profile)


def test_excel_path_profile_preserves_header_offset_and_types(tmp_path):
    path = tmp_path / "example.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Report"])
    sheet.append(["id", "amount"])
    for index in range(2000):
        sheet.append([str(index), index * 0.5])
    workbook.save(path)
    _, profile = ingestion.build_profile_path(path, path.name)
    table = profile["tables"][0]
    assert table["header_row_index"] == 1
    assert [column["name"] for column in table["columns"]] == ["id", "amount"]
    assert table["record_count"] is None


def test_dimensionless_excel_does_not_scan_sheet_before_sampling(tmp_path, monkeypatch):
    from openpyxl.worksheet._reader import WorkSheetParser
    from app.services.excel_profile_reader import load_profile_workbook

    path = tmp_path / "dimensionless.xlsx"
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet()
    sheet.append(["id", "date"])
    for index in range(4000):
        sheet.append([index, datetime(2026, 1, 2)])
    workbook.save(path)
    with ZipFile(path) as archive:
        assert b'<dimension ' not in archive.read('xl/worksheets/sheet1.xml')
    consumed = 0
    original = WorkSheetParser.parse_row

    def count_rows(self, element):
        nonlocal consumed
        consumed += 1
        return original(self, element)

    monkeypatch.setattr(WorkSheetParser, "parse_row", count_rows)
    reader = load_profile_workbook(path, row_limit=ingestion.MAX_PROFILE_ROWS + 26,
                                   column_limit=ingestion.MAX_PROFILE_COLUMNS)
    try:
        assert consumed == 0
        profile = ingestion._sheet_profile_rows(reader.active.iter_rows(values_only=True))
        assert consumed <= ingestion.MAX_PROFILE_ROWS + 26
        assert profile[2][0][1] == datetime(2026, 1, 2)
    finally:
        reader.close()
